#!/usr/bin/env python3
"""
compare_demosaic.py
对比同一 rawengine-cli 使用不同 demosaic 算法解码的输出质量和速度。

用法示例:
  python compare_demosaic.py
  python compare_demosaic.py --dir-a amazebilinear --dir-b rcd

默认比较:
  C:/pack/file/output/rawtherapee_amazebilinear  vs  C:/pack/file/output/rawtherapee_rcd

质量指标: PSNR, MAE(R/G/B), delta_luma_mean
速度指标: 读取各目录 *_time.txt 中的耗时记录
输出: compare_demosaic_report.xlsx（与本脚本同目录）
"""

from __future__ import annotations

import argparse
import math
from pathlib import Path, PurePosixPath
from typing import Dict, List

import numpy as np
from PIL import Image

try:
    from openpyxl import Workbook
except Exception as exc:
    raise SystemExit(f"缺少 openpyxl: pip install openpyxl\n{exc}")

OUTPUT_BASE = r"C:\pack\file\output"
IMAGE_EXTS  = {".jpg", ".jpeg", ".png", ".tif", ".tiff"}


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="对比两种 demosaic 算法的解码质量和速度")
    parser.add_argument("--dir-a", default="amazebilinear",
                        help="算法 A 名称 (输出目录后缀，默认 amazebilinear)")
    parser.add_argument("--dir-b", default="rcd",
                        help="算法 B 名称 (输出目录后缀，默认 rcd)")
    parser.add_argument("--base", default=OUTPUT_BASE,
                        help="输出根目录，两个算法目录在此下以 rawtherapee_<name> 命名")
    parser.add_argument("--out",
                        default=str(Path(__file__).resolve().parent / "compare_demosaic_report.xlsx"),
                        help="输出 Excel 文件路径")
    return parser.parse_args()


def norm_key(rel: str) -> str:
    """归一化匹配键: 父目录/文件名主干（小写、正斜杠），忽略扩展名差异。
    *_time.txt 记录的是输入名(.CR3/.ARW)，图片索引是输出名(.jpg)，需按 stem 匹配。"""
    p = PurePosixPath(str(rel).replace("\\", "/"))
    parent = str(p.parent)
    stem = p.stem
    key = stem if parent in (".", "") else f"{parent}/{stem}"
    return key.lower()


def list_images(root: Path) -> Dict[str, Path]:
    return {
        norm_key(str(p.relative_to(root))): p
        for p in root.rglob("*")
        if p.is_file() and p.suffix.lower() in IMAGE_EXTS
    }


def load_rgb(path: Path) -> np.ndarray:
    with Image.open(path) as img:
        return np.asarray(img.convert("RGB"), dtype=np.float32)


def compute_metrics(a_img: np.ndarray, b_img: np.ndarray) -> Dict[str, float]:
    """计算两张图片之间的差异指标（a=算法A，b=算法B）"""
    if a_img.shape != b_img.shape:
        h = min(a_img.shape[0], b_img.shape[0])
        w = min(a_img.shape[1], b_img.shape[1])
        a_img = a_img[:h, :w, :]
        b_img = b_img[:h, :w, :]

    diff = a_img.astype(np.float64) - b_img.astype(np.float64)
    abs_diff = np.abs(diff)

    a_luma = 0.2126 * a_img[:, :, 0] + 0.7152 * a_img[:, :, 1] + 0.0722 * a_img[:, :, 2]
    b_luma = 0.2126 * b_img[:, :, 0] + 0.7152 * b_img[:, :, 1] + 0.0722 * b_img[:, :, 2]

    mse = float(np.mean(diff * diff))
    psnr = 99.0 if mse <= 1e-12 else 20.0 * math.log10(255.0 / math.sqrt(mse))

    return {
        "a_luma_mean": float(a_luma.mean()),
        "b_luma_mean": float(b_luma.mean()),
        "delta_luma": float(a_luma.mean() - b_luma.mean()),
        "mae_r": float(abs_diff[:, :, 0].mean()),
        "mae_g": float(abs_diff[:, :, 1].mean()),
        "mae_b": float(abs_diff[:, :, 2].mean()),
        "mae_rgb": float(abs_diff.mean()),
        "psnr": float(psnr),
        "width": a_img.shape[1],
        "height": a_img.shape[0],
    }


def read_time_files(root: Path) -> Dict[str, float]:
    """读取所有 *_time.txt，返回 {相对路径(小写): 耗时秒} 映射"""
    result: Dict[str, float] = {}
    for f in root.rglob("*_time.txt"):
        try:
            for line in f.read_text(encoding="utf-8").splitlines():
                line = line.strip()
                if not line.startswith("-"):
                    continue
                # 格式: - relative/path.cr2: 12.345678 [OK]
                line = line.lstrip("- ").strip()
                colon_pos = line.rfind(":")
                if colon_pos < 0:
                    continue
                rel = line[:colon_pos].strip()
                rest = line[colon_pos + 1:].strip()
                time_str = rest.split()[0]
                result[norm_key(rel)] = float(time_str)
        except Exception:
            continue
    return result


def safe_sheet_name(name: str, used: set) -> str:
    cleaned = (name.replace("\\", "_").replace("/", "_")
                   .replace(":", "_").replace("*", "_").replace("?", "_")
                   .replace("[", "(").replace("]", ")")) or "root"
    base = cleaned[:31]
    candidate, idx = base, 1
    while candidate in used:
        suffix = f"_{idx}"
        candidate = base[:31 - len(suffix)] + suffix
        idx += 1
    used.add(candidate)
    return candidate


def write_excel(rows: List[Dict], dir_a_name: str, dir_b_name: str, out_file: Path) -> None:
    wb = Workbook()
    wb.remove(wb.active)

    headers = [
        "relative_path",
        f"{dir_a_name}_path",
        f"{dir_b_name}_path",
        f"{dir_a_name}_time_s",
        f"{dir_b_name}_time_s",
        "speedup(A/B)",
        f"{dir_a_name}_luma",
        f"{dir_b_name}_luma",
        "delta_luma(A-B)",
        "mae_r", "mae_g", "mae_b", "mae_rgb",
        "psnr(A_vs_B)",
        "width", "height",
    ]

    by_group: Dict[str, List[Dict]] = {}
    for r in rows:
        parent = str(Path(str(r["relative_path"])).parent).replace("\\", "/")
        if parent == ".":
            parent = "root"
        by_group.setdefault(parent, []).append(r)

    used: set = set()
    summary = wb.create_sheet(title="summary")
    summary.append([
        "group", "images_compared",
        f"avg_{dir_a_name}_time_s", f"avg_{dir_b_name}_time_s",
        "avg_speedup(A/B)",
        "avg_psnr", "avg_mae_rgb", "avg_abs_delta_luma",
    ])

    for group, group_rows in sorted(by_group.items()):
        ws = wb.create_sheet(title=safe_sheet_name(group, used))
        ws.append(headers)

        times_a, times_b, speedups = [], [], []
        psnrs, maes, delta_lumas = [], [], []

        for r in group_rows:
            ws.append([r.get(h, "") for h in headers])
            if isinstance(r.get(f"{dir_a_name}_time_s"), float) and isinstance(r.get(f"{dir_b_name}_time_s"), float):
                ta, tb = r[f"{dir_a_name}_time_s"], r[f"{dir_b_name}_time_s"]
                times_a.append(ta)
                times_b.append(tb)
                if tb > 0:
                    speedups.append(ta / tb)
            if isinstance(r.get("psnr(A_vs_B)"), float):
                psnrs.append(r["psnr(A_vs_B)"])
            if isinstance(r.get("mae_rgb"), float):
                maes.append(r["mae_rgb"])
            if isinstance(r.get("delta_luma(A-B)"), float):
                delta_lumas.append(abs(r["delta_luma(A-B)"]))

        def avg(lst): return round(float(np.mean(lst)), 4) if lst else ""

        summary.append([
            group,
            len(group_rows),
            avg(times_a), avg(times_b),
            avg(speedups),
            avg(psnrs), avg(maes), avg(delta_lumas),
        ])

    out_file.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_file)
    print(f"报告已写入: {out_file}")


def main() -> None:
    args = parse_args()

    dir_a = Path(args.base) / f"rawtherapee_{args.dir_a}"
    dir_b = Path(args.base) / f"rawtherapee_{args.dir_b}"

    if not dir_a.is_dir():
        raise SystemExit(f"算法 A 输出目录不存在: {dir_a}\n"
                         f"请先运行: python test_rawengine.py --method {args.dir_a}")
    if not dir_b.is_dir():
        raise SystemExit(f"算法 B 输出目录不存在: {dir_b}\n"
                         f"请先运行: python test_rawengine.py --method {args.dir_b}")

    print(f"算法 A: {args.dir_a}  →  {dir_a}")
    print(f"算法 B: {args.dir_b}  →  {dir_b}")

    map_a = list_images(dir_a)
    map_b = list_images(dir_b)
    times_a = read_time_files(dir_a)
    times_b = read_time_files(dir_b)

    common = sorted(set(map_a.keys()) & set(map_b.keys()))
    if not common:
        raise SystemExit("两个目录没有可匹配的同名图片，请先完成两边解码输出。")

    print(f"共匹配 {len(common)} 张图片，开始计算质量指标...")

    rows: List[Dict] = []
    for idx, rel in enumerate(common, 1):
        print(f"  [{idx}/{len(common)}] {rel}")
        ta = times_a.get(rel)
        tb = times_b.get(rel)
        speedup = round(ta / tb, 4) if ta and tb and tb > 0 else ""

        try:
            a_img = load_rgb(map_a[rel])
            b_img = load_rgb(map_b[rel])
            m = compute_metrics(a_img, b_img)
            row = {
                "relative_path": rel,
                f"{args.dir_a}_path": str(map_a[rel]),
                f"{args.dir_b}_path": str(map_b[rel]),
                f"{args.dir_a}_time_s": ta,
                f"{args.dir_b}_time_s": tb,
                "speedup(A/B)": speedup,
                f"{args.dir_a}_luma": round(m["a_luma_mean"], 4),
                f"{args.dir_b}_luma": round(m["b_luma_mean"], 4),
                "delta_luma(A-B)": round(m["delta_luma"], 4),
                "mae_r": round(m["mae_r"], 4),
                "mae_g": round(m["mae_g"], 4),
                "mae_b": round(m["mae_b"], 4),
                "mae_rgb": round(m["mae_rgb"], 4),
                "psnr(A_vs_B)": round(m["psnr"], 4),
                "width": m["width"],
                "height": m["height"],
            }
        except Exception as exc:
            row = {
                "relative_path": rel,
                f"{args.dir_a}_path": str(map_a[rel]),
                f"{args.dir_b}_path": str(map_b[rel]),
                f"{args.dir_a}_time_s": ta,
                f"{args.dir_b}_time_s": tb,
                "speedup(A/B)": speedup,
                "psnr(A_vs_B)": f"ERROR: {exc}",
            }
        rows.append(row)

    write_excel(rows, args.dir_a, args.dir_b, Path(args.out))
    print(f"\n共对比 {len(rows)} 张图片。")
    print("查看 summary sheet 获取速度/质量汇总。")


if __name__ == "__main__":
    main()
